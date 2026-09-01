from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from ea_node_editor.addons.ansys_dpf.property_edit_adapter import AnsysDpfPropertyEditAdapter
from ea_node_editor.graph.model import GraphModel
from ea_node_editor.nodes.bootstrap import build_default_registry
from ea_node_editor.ui.shell.controllers.workspace_view_nav_ops import WorkspaceViewNavOps
from ea_node_editor.ui.shell.inspector_projection import build_selected_node_property_items
from ea_node_editor.ui_qml.graph_scene_bridge import GraphSceneBridge
from tests.passive_property_editor_fixtures import (
    PASSIVE_EDITOR_FIXTURE_TYPE_ID,
    register_passive_editor_fixture,
)


class PassivePropertyEditorModeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.registry = build_default_registry()
        register_passive_editor_fixture(self.registry)
        self.model = GraphModel()
        self.workspace_id = self.model.active_workspace.workspace_id
        self.scene = GraphSceneBridge()
        self.scene.set_workspace(self.model, self.registry, self.workspace_id)

    def _property_items_for_type(self, type_id: str) -> dict[str, dict[str, object]]:
        node_id = self.scene.add_node_from_type(type_id, 0.0, 0.0)
        workspace = self.model.project.workspaces[self.workspace_id]
        node = workspace.nodes[node_id]
        spec = self.registry.get_spec(type_id)
        items = build_selected_node_property_items(
            node=node,
            spec=spec,
            subnode_pin_type_ids=set(),
        )
        return {str(item["key"]): item for item in items}

    def _property_items_for_node(
        self,
        node_id: str,
        *,
        project_path: str | None = None,
        dpf_named_selection_provider=None,
        dpf_result_field_options_provider=None,
    ) -> dict[str, dict[str, object]]:
        workspace = self.model.project.workspaces[self.workspace_id]
        node = workspace.nodes[node_id]
        spec = self.registry.get_spec(node.type_id)
        property_edit_adapters = None
        if dpf_named_selection_provider is not None or dpf_result_field_options_provider is not None:
            def metadata_provider(result_path: str):  # noqa: ANN202
                options = (
                    dict(dpf_result_field_options_provider(result_path))
                    if dpf_result_field_options_provider is not None
                    else {}
                )
                if dpf_named_selection_provider is not None:
                    options["named_selection"] = tuple(dpf_named_selection_provider(result_path))
                return options

            property_edit_adapters = (AnsysDpfPropertyEditAdapter(metadata_provider=metadata_provider),)
        items = build_selected_node_property_items(
            node=node,
            spec=spec,
            subnode_pin_type_ids=set(),
            workspace_nodes=workspace.nodes,
            workspace_edges=workspace.edges,
            project_path=project_path,
            property_edit_adapters=property_edit_adapters,
        )
        return {str(item["key"]): item for item in items}

    def test_sdk_driven_editor_modes_follow_property_metadata_not_property_keys(self) -> None:
        items = self._property_items_for_type(PASSIVE_EDITOR_FIXTURE_TYPE_ID)

        self.assertEqual(items["notes_blob"]["editor_mode"], "textarea")
        self.assertEqual(items["media_ref"]["editor_mode"], "path")
        self.assertEqual(items["accent_color"]["editor_mode"], "color")
        self.assertEqual(items["caption"]["editor_mode"], "text")

    def test_existing_logger_property_editor_modes_remain_text_and_enum(self) -> None:
        items = self._property_items_for_type("core.logger")

        self.assertEqual(items["message"]["editor_mode"], "text")
        self.assertEqual(items["level"]["editor_mode"], "enum")

    def test_planning_family_long_form_fields_use_textarea(self) -> None:
        task_items = self._property_items_for_type("passive.planning.task_card")
        risk_items = self._property_items_for_type("passive.planning.risk_card")
        decision_items = self._property_items_for_type("passive.planning.decision_card")

        self.assertEqual(task_items["body"]["editor_mode"], "textarea")
        self.assertEqual(risk_items["body"]["editor_mode"], "textarea")
        self.assertEqual(risk_items["mitigation"]["editor_mode"], "textarea")
        self.assertEqual(decision_items["body"]["editor_mode"], "textarea")
        self.assertEqual(decision_items["outcome"]["editor_mode"], "textarea")
        self.assertEqual(task_items["owner"]["editor_mode"], "text")
        self.assertEqual(task_items["status"]["editor_mode"], "enum")

    def test_annotation_family_uses_textarea_only_for_body_fields(self) -> None:
        sticky_items = self._property_items_for_type("passive.annotation.sticky_note")
        callout_items = self._property_items_for_type("passive.annotation.callout")
        section_items = self._property_items_for_type("passive.annotation.section_header")
        text_items = self._property_items_for_type("passive.annotation.text")

        self.assertEqual(sticky_items["body"]["editor_mode"], "textarea")
        self.assertEqual(callout_items["body"]["editor_mode"], "textarea")
        self.assertEqual(text_items["text"]["editor_mode"], "textarea")
        self.assertEqual(text_items["text_color"]["editor_mode"], "color")
        self.assertEqual(text_items["background_color"]["editor_mode"], "color")
        self.assertEqual(sticky_items["title"]["editor_mode"], "text")
        self.assertEqual(section_items["title"]["editor_mode"], "text")
        self.assertEqual(section_items["subtitle"]["editor_mode"], "text")

    def test_flowchart_family_exposes_text_title_editor(self) -> None:
        decision_items = self._property_items_for_type("passive.flowchart.decision")
        database_items = self._property_items_for_type("passive.flowchart.database")
        decision_spec = self.registry.get_spec("passive.flowchart.decision")
        decision_body = next(prop for prop in decision_spec.properties if prop.key == "body")

        self.assertEqual(decision_items["title"]["editor_mode"], "text")
        self.assertEqual(decision_items["title"]["value"], "Decision")
        self.assertEqual(decision_items["body"]["editor_mode"], "textarea")
        self.assertEqual(decision_items["body"]["value"], "Decision")
        self.assertEqual(database_items["title"]["editor_mode"], "text")
        self.assertEqual(database_items["title"]["value"], "Database")
        self.assertEqual(database_items["body"]["editor_mode"], "textarea")
        self.assertEqual(database_items["body"]["value"], "Database")
        self.assertTrue(WorkspaceViewNavOps._property_is_content_searchable(decision_body))

    def test_dpf_mesh_scoping_named_selection_uses_editable_combo_with_upstream_options(self) -> None:
        result_file_id = self.scene.add_node_from_type("dpf.result_file", 0.0, 0.0)
        model_id = self.scene.add_node_from_type("dpf.model", 320.0, 0.0)
        scoping_id = self.scene.add_node_from_type("dpf.scoping.mesh", 640.0, 0.0)

        self.scene.set_node_property(result_file_id, "path", "C:/tmp/demo.rst")
        self.scene.add_edge(result_file_id, "result_file", model_id, "result_file")
        self.scene.add_edge(model_id, "model", scoping_id, "model")

        requested_paths: list[str] = []

        def named_selection_provider(result_path: str) -> tuple[str, ...]:
            requested_paths.append(result_path)
            return ("BOLT_NODES", "FACE_A")

        items = self._property_items_for_node(
            scoping_id,
            dpf_named_selection_provider=named_selection_provider,
        )

        self.assertEqual(requested_paths, ["C:\\tmp\\demo.rst"])
        self.assertEqual(items["named_selection"]["editor_mode"], "editable_combo")
        self.assertEqual(items["named_selection"]["enum_values"], ["BOLT_NODES", "FACE_A"])
        self.assertEqual(items["named_selection"].get("placeholder_text"), "")

    def test_dpf_mesh_scoping_named_selection_shows_empty_state_when_upstream_model_has_none(self) -> None:
        model_id = self.scene.add_node_from_type("dpf.model", 320.0, 0.0)
        scoping_id = self.scene.add_node_from_type("dpf.scoping.mesh", 640.0, 0.0)

        self.scene.set_node_property(model_id, "path", "C:/tmp/demo.rst")
        self.scene.add_edge(model_id, "model", scoping_id, "model")

        items = self._property_items_for_node(
            scoping_id,
            dpf_named_selection_provider=lambda result_path: (),
        )

        self.assertEqual(items["named_selection"]["editor_mode"], "editable_combo")
        self.assertEqual(items["named_selection"]["enum_values"], [])
        self.assertEqual(items["named_selection"].get("placeholder_text"), "No named selections found")

    def test_dpf_mesh_scoping_named_selection_resolves_relative_model_path_against_project(self) -> None:
        model_id = self.scene.add_node_from_type("dpf.model", 320.0, 0.0)
        scoping_id = self.scene.add_node_from_type("dpf.scoping.mesh", 640.0, 0.0)

        self.scene.set_node_property(model_id, "path", "results/demo.rst")
        self.scene.add_edge(model_id, "model", scoping_id, "model")

        requested_paths: list[str] = []

        def named_selection_provider(result_path: str) -> tuple[str, ...]:
            requested_paths.append(result_path)
            return ("BOLT_NODES",)

        project_path = str(Path("C:/workflows/project.json"))
        items = self._property_items_for_node(
            scoping_id,
            project_path=project_path,
            dpf_named_selection_provider=named_selection_provider,
        )

        self.assertEqual(
            requested_paths,
            [str(Path("C:/workflows/results/demo.rst").resolve(strict=False))],
        )
        self.assertEqual(items["named_selection"]["editor_mode"], "editable_combo")
        self.assertEqual(items["named_selection"]["enum_values"], ["BOLT_NODES"])

    def test_dpf_mesh_scoping_non_named_selection_mode_keeps_manual_text_editor(self) -> None:
        scoping_id = self.scene.add_node_from_type("dpf.scoping.mesh", 640.0, 0.0)
        self.scene.set_node_property(scoping_id, "selection_mode", "node_ids")

        items = self._property_items_for_node(scoping_id)

        self.assertEqual(items["named_selection"]["editor_mode"], "text")
        self.assertEqual(items["named_selection"]["enum_values"], [])

    def test_dpf_result_field_selection_values_use_editable_combos_with_upstream_options(self) -> None:
        result_file_id = self.scene.add_node_from_type("dpf.result_file", 0.0, 0.0)
        model_id = self.scene.add_node_from_type("dpf.model", 320.0, 0.0)
        result_field_id = self.scene.add_node_from_type("dpf.result_field", 640.0, 0.0)

        self.scene.set_node_property(result_file_id, "path", "C:/tmp/demo.rst")
        self.scene.add_edge(result_file_id, "result_file", model_id, "result_file")
        self.scene.add_edge(model_id, "model", result_field_id, "model")

        requested_paths: list[str] = []

        def result_field_options_provider(result_path: str) -> dict[str, tuple[str, ...]]:
            requested_paths.append(result_path)
            return {
                "result_name": ("displacement", "stress", "Displacement"),
                "set_ids": ("1", "2", "2"),
                "time_values": ("0.0", "1.25"),
            }

        items = self._property_items_for_node(
            result_field_id,
            dpf_result_field_options_provider=result_field_options_provider,
        )

        self.assertEqual(requested_paths, ["C:\\tmp\\demo.rst"])
        self.assertEqual(items["result_name"]["editor_mode"], "editable_combo")
        self.assertEqual(items["result_name"]["enum_values"], ["displacement", "stress"])
        self.assertEqual(items["result_name"].get("placeholder_text"), "")
        self.assertEqual(items["set_ids"]["editor_mode"], "editable_combo")
        self.assertEqual(items["set_ids"]["enum_values"], ["1", "2"])
        self.assertEqual(items["set_ids"].get("placeholder_text"), "")
        self.assertEqual(items["time_values"]["editor_mode"], "editable_combo")
        self.assertEqual(items["time_values"]["enum_values"], ["0.0", "1.25"])
        self.assertEqual(items["time_values"].get("placeholder_text"), "")

    def test_dpf_workflow_result_fields_resolves_options_from_result_source_model(self) -> None:
        source_id = self.scene.add_node_from_type("dpf.workflow.result_source", 0.0, 0.0)
        result_fields_id = self.scene.add_node_from_type("dpf.workflow.result_fields", 320.0, 0.0)

        self.scene.set_node_property(source_id, "path", "C:/tmp/demo.rst")
        self.scene.set_node_property(result_fields_id, "time_scope_mode", "set_ids")
        self.scene.add_edge(source_id, "model", result_fields_id, "model")

        requested_paths: list[str] = []

        def result_field_options_provider(result_path: str) -> dict[str, tuple[str, ...]]:
            requested_paths.append(result_path)
            return {
                "result_name": ("displacement", "stress"),
                "set_ids": ("1", "2"),
                "time_values": ("0.0", "1.25"),
            }

        items = self._property_items_for_node(
            result_fields_id,
            dpf_result_field_options_provider=result_field_options_provider,
        )

        self.assertEqual(requested_paths, ["C:\\tmp\\demo.rst"])
        self.assertEqual(items["result_name"]["editor_mode"], "editable_combo")
        self.assertEqual(items["result_name"]["enum_values"], ["displacement", "stress"])
        self.assertEqual(items["set_ids"]["editor_mode"], "chip_list")
        self.assertEqual(items["set_ids"]["enum_values"], ["1", "2"])
        self.assertNotIn("time_values", items)

    def test_dpf_workflow_result_viewer_resolves_options_from_own_path_property(self) -> None:
        viewer_id = self.scene.add_node_from_type("dpf.workflow.result_viewer", 0.0, 0.0)
        self.scene.set_node_property(viewer_id, "path", "C:/tmp/demo.rst")
        self.scene.set_node_property(viewer_id, "selection_mode", "named_selection")
        self.scene.set_node_property(viewer_id, "time_scope_mode", "set_ids")

        result_paths: list[str] = []
        named_selection_paths: list[str] = []

        def result_field_options_provider(result_path: str) -> dict[str, tuple[str, ...]]:
            result_paths.append(result_path)
            return {
                "result_name": ("displacement",),
                "set_ids": ("1",),
                "time_values": ("0.0",),
            }

        def named_selection_provider(result_path: str) -> tuple[str, ...]:
            named_selection_paths.append(result_path)
            return ("BOLT_NODES",)

        items = self._property_items_for_node(
            viewer_id,
            dpf_named_selection_provider=named_selection_provider,
            dpf_result_field_options_provider=result_field_options_provider,
        )

        self.assertEqual(result_paths, ["C:\\tmp\\demo.rst"])
        self.assertEqual(named_selection_paths, ["C:\\tmp\\demo.rst"])
        self.assertEqual(items["result_name"]["editor_mode"], "editable_combo")
        self.assertEqual(items["result_name"]["enum_values"], ["displacement"])
        self.assertEqual(items["set_ids"]["editor_mode"], "chip_list")
        self.assertEqual(items["set_ids"]["enum_values"], ["1"])
        self.assertNotIn("time_values", items)
        self.assertEqual(items["named_selection"]["editor_mode"], "editable_combo")
        self.assertEqual(items["named_selection"]["enum_values"], ["BOLT_NODES"])

    def test_dpf_workflow_post_nodes_resolve_options_from_model_input(self) -> None:
        for node_type_id in (
            "dpf.workflow.min_max_envelope",
            "dpf.workflow.time_history_probe",
            "dpf.workflow.stress_invariants",
        ):
            with self.subTest(node_type_id=node_type_id):
                source_id = self.scene.add_node_from_type("dpf.workflow.result_source", 0.0, 0.0)
                node_id = self.scene.add_node_from_type(node_type_id, 320.0, 0.0)
                self.scene.set_node_property(source_id, "path", "C:/tmp/demo.rst")
                self.scene.set_node_property(node_id, "selection_mode", "named_selection")
                self.scene.set_node_property(node_id, "time_scope_mode", "set_ids")
                self.scene.add_edge(source_id, "model", node_id, "model")

                def result_field_options_provider(result_path: str) -> dict[str, tuple[str, ...]]:
                    return {
                        "result_name": ("displacement", "stress"),
                        "set_ids": ("1", "2"),
                        "time_values": ("0.0", "1.25"),
                    }

                items = self._property_items_for_node(
                    node_id,
                    dpf_named_selection_provider=lambda result_path: ("BOLT_NODES",),
                    dpf_result_field_options_provider=result_field_options_provider,
                )

                self.assertEqual(items["set_ids"]["editor_mode"], "chip_list")
                self.assertEqual(items["set_ids"]["enum_values"], ["1", "2"])
                self.assertEqual(items["named_selection"]["editor_mode"], "editable_combo")
                self.assertEqual(items["named_selection"]["enum_values"], ["BOLT_NODES"])
                if node_type_id != "dpf.workflow.stress_invariants":
                    self.assertEqual(items["result_name"]["editor_mode"], "editable_combo")
                    self.assertEqual(
                        items["result_name"]["enum_values"], ["displacement", "stress"]
                    )

    def test_dpf_mode_shape_viewer_mode_options_mirror_set_ids_from_own_path(self) -> None:
        viewer_id = self.scene.add_node_from_type("dpf.workflow.mode_shape_viewer", 0.0, 0.0)
        self.scene.set_node_property(viewer_id, "path", "C:/tmp/modal.rst")

        def result_field_options_provider(result_path: str) -> dict[str, tuple[str, ...]]:
            return {
                "result_name": ("displacement",),
                "set_ids": ("1", "2", "3", "4", "5", "6", "7"),
                "time_values": ("2118.5", "2241.2"),
            }

        items = self._property_items_for_node(
            viewer_id,
            dpf_result_field_options_provider=result_field_options_provider,
        )

        self.assertEqual(items["mode"]["editor_mode"], "editable_combo")
        self.assertEqual(items["mode"]["enum_values"], ["1", "2", "3", "4", "5", "6", "7"])
        self.assertEqual(items["mode"].get("placeholder_text"), "")

    def test_dpf_mode_shape_viewer_mode_placeholders_cover_missing_path_and_empty_model(self) -> None:
        unconnected_id = self.scene.add_node_from_type("dpf.workflow.mode_shape_viewer", 0.0, 0.0)
        items = self._property_items_for_node(
            unconnected_id,
            dpf_result_field_options_provider=lambda result_path: {},
        )
        self.assertEqual(items["mode"]["editor_mode"], "editable_combo")
        self.assertEqual(items["mode"]["enum_values"], [])
        self.assertEqual(
            items["mode"].get("placeholder_text"), "Set a modal result file to load modes"
        )

        empty_model_id = self.scene.add_node_from_type("dpf.workflow.mode_shape_viewer", 320.0, 0.0)
        self.scene.set_node_property(empty_model_id, "path", "C:/tmp/modal.rst")
        items = self._property_items_for_node(
            empty_model_id,
            dpf_result_field_options_provider=lambda result_path: {},
        )
        self.assertEqual(items["mode"]["enum_values"], [])
        self.assertEqual(items["mode"].get("placeholder_text"), "No modes found")

    def test_dpf_result_field_selection_values_show_empty_states_when_model_has_none(self) -> None:
        model_id = self.scene.add_node_from_type("dpf.model", 320.0, 0.0)
        result_field_id = self.scene.add_node_from_type("dpf.result_field", 640.0, 0.0)

        self.scene.set_node_property(model_id, "path", "C:/tmp/demo.rst")
        self.scene.add_edge(model_id, "model", result_field_id, "model")

        items = self._property_items_for_node(
            result_field_id,
            dpf_result_field_options_provider=lambda result_path: {},
        )

        self.assertEqual(items["result_name"]["editor_mode"], "editable_combo")
        self.assertEqual(
            items["result_name"]["enum_values"],
            ["displacement", "stress", "elastic_strain", "structural_temperature"],
        )
        self.assertEqual(items["result_name"].get("placeholder_text"), "No result names found")
        self.assertEqual(items["set_ids"]["editor_mode"], "editable_combo")
        self.assertEqual(items["set_ids"]["enum_values"], [])
        self.assertEqual(items["set_ids"].get("placeholder_text"), "No set IDs found")
        self.assertEqual(items["time_values"]["editor_mode"], "editable_combo")
        self.assertEqual(items["time_values"]["enum_values"], [])
        self.assertEqual(items["time_values"].get("placeholder_text"), "No time values found")

    def test_dpf_result_field_selection_values_stay_editable_without_upstream_model(self) -> None:
        result_field_id = self.scene.add_node_from_type("dpf.result_field", 640.0, 0.0)

        def result_field_options_provider(result_path: str) -> dict[str, tuple[str, ...]]:
            raise AssertionError(f"provider should not be called without a model path: {result_path}")

        items = self._property_items_for_node(
            result_field_id,
            dpf_result_field_options_provider=result_field_options_provider,
        )

        self.assertEqual(items["result_name"]["editor_mode"], "editable_combo")
        self.assertEqual(
            items["result_name"]["enum_values"],
            ["displacement", "stress", "elastic_strain", "structural_temperature"],
        )
        self.assertEqual(
            items["result_name"].get("placeholder_text"),
            "Connect a DPF model to load result names",
        )
        self.assertEqual(items["set_ids"]["editor_mode"], "editable_combo")
        self.assertEqual(items["set_ids"]["enum_values"], [])
        self.assertEqual(
            items["set_ids"].get("placeholder_text"),
            "Connect a DPF model to load set IDs",
        )
        self.assertEqual(items["time_values"]["editor_mode"], "editable_combo")
        self.assertEqual(items["time_values"]["enum_values"], [])
        self.assertEqual(
            items["time_values"].get("placeholder_text"),
            "Connect a DPF model to load time values",
        )

    def test_tabular_data_source_uses_editable_combo_with_scanned_objects(self) -> None:
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        if self.registry.spec_or_none("tabular.input") is None:
            self.skipTest("tabular.input add-on is not registered in the default registry")

        with tempfile.TemporaryDirectory() as tmp_dir:
            source = Path(tmp_dir) / "book.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "First"
            workbook["First"].append(["name", "value"])
            second = workbook.create_sheet("Second")
            second.append(["name", "value"])
            workbook.save(source)

            node_id = self.scene.add_node_from_type("tabular.input", 0.0, 0.0)
            self.scene.set_node_property(node_id, "path", str(source))
            items = self._property_items_for_node(node_id)

        self.assertEqual(items["selected_object"]["editor_mode"], "editable_combo")
        self.assertEqual(items["selected_object"]["enum_values"], ["First", "Second"])
        self.assertEqual(items["selected_object"].get("placeholder_text"), "Select a sheet, key, or dataset")

    def test_tabular_data_source_selector_remains_editable_without_a_path(self) -> None:
        if self.registry.spec_or_none("tabular.input") is None:
            self.skipTest("tabular.input add-on is not registered in the default registry")

        node_id = self.scene.add_node_from_type("tabular.input", 0.0, 0.0)
        items = self._property_items_for_node(node_id)

        self.assertEqual(items["selected_object"]["editor_mode"], "editable_combo")
        self.assertEqual(items["selected_object"]["enum_values"], [])
        self.assertEqual(items["selected_object"].get("placeholder_text"), "Choose a tabular data file first")


if __name__ == "__main__":
    unittest.main()
